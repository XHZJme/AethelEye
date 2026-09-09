"""
筱和灵眸(AethelEye) - 数据导出API

Excel导出异动记录、价格记录等
"""
from datetime import datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import io

from app.database import get_db_session
from app.models.alert import Alert
from app.models.task import MonitorTask
from app.models.sku import SKU
from app.models.price_record import PriceRecord
from app.models.user import User
from app.middleware.permission import require_operator
from app.utils.logger import get_logger

# 日志
log = get_logger("audit")

# 路由
router = APIRouter(prefix="/api/export", tags=["数据导出"])


@router.get("/alerts")
def export_alerts(
    status: Optional[str] = None,
    platform: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    session: Session = Depends(get_db_session),
    user: User = Depends(require_operator)
):
    """
    导出异动记录为Excel

    使用openpyxl生成Excel文件（如果没有可用csv格式备选）
    """
    try:
        # 尝试使用openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        # 查询数据
        query = session.query(Alert)
        if platform:
            query = query.join(MonitorTask, Alert.task_id == MonitorTask.id).filter(MonitorTask.platform == platform)

        if status:
            query = query.filter(Alert.status == status)
        if start_date:
            query = query.filter(Alert.detected_at >= start_date)
        if end_date:
            query = query.filter(Alert.detected_at <= end_date)

        alerts = query.order_by(Alert.detected_at.desc()).all()

        task_ids = sorted({a.task_id for a in alerts if a.task_id})
        sku_ids = sorted({a.sku_id for a in alerts if a.sku_id})
        task_map = {}
        sku_map = {}
        if task_ids:
            tasks = session.query(MonitorTask).filter(MonitorTask.id.in_(task_ids)).all()
            task_map = {t.id: t for t in tasks}
        if sku_ids:
            skus = session.query(SKU).filter(SKU.id.in_(sku_ids)).all()
            sku_map = {s.id: s for s in skus}

        # 创建Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "异动记录"

        # 标题行
        headers = ["序号", "发现时间", "店铺", "商品", "SKU", "平台", "基准价", "异动到手价", "价差", "价差率", "状态", "截图"]
        ws.append(headers)

        # 设置标题样式
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        # 数据行
        for i, alert in enumerate(alerts, 1):
            task = task_map.get(alert.task_id)
            sku = sku_map.get(alert.sku_id)

            status_map = {
                "new": "未处理",
                "confirmed": "已确认",
                "resolved": "已处理",
                "false_alarm": "误报",
            }

            row = [
                i,
                alert.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
                task.shop_name if task else "[任务已删除]",
                task.product_title if task else "[商品信息缺失]",
                sku.sku_name if sku else "[SKU信息缺失]",
                task.platform if task else "未知",
                f"¥{alert.base_price:.2f}",
                f"¥{alert.alert_price:.2f}",
                f"¥{abs(alert.price_diff):.2f}",
                f"{abs(alert.price_diff_pct):.1f}%",
                status_map.get(alert.status, alert.status),
                alert.screenshot_path or "",
            ]
            ws.append(row)

        # 保存到内存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        log.info(
            f"异动记录已导出",
            data={
                "exported_by": user.username,
                "record_count": len(alerts),
                "format": "excel",
            }
        )

        # 返回文件
        filename = f"alerts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except ImportError:
        # openpyxl不可用，使用CSV格式
        import csv

        query = session.query(Alert)
        if platform:
            query = query.join(MonitorTask, Alert.task_id == MonitorTask.id).filter(MonitorTask.platform == platform)

        if status:
            query = query.filter(Alert.status == status)
        if start_date:
            query = query.filter(Alert.detected_at >= start_date)
        if end_date:
            query = query.filter(Alert.detected_at <= end_date)

        alerts = query.order_by(Alert.detected_at.desc()).all()

        task_ids = sorted({a.task_id for a in alerts if a.task_id})
        sku_ids = sorted({a.sku_id for a in alerts if a.sku_id})
        task_map = {}
        sku_map = {}
        if task_ids:
            tasks = session.query(MonitorTask).filter(MonitorTask.id.in_(task_ids)).all()
            task_map = {t.id: t for t in tasks}
        if sku_ids:
            skus = session.query(SKU).filter(SKU.id.in_(sku_ids)).all()
            sku_map = {s.id: s for s in skus}

        output = io.StringIO()
        writer = csv.writer(output)

        headers = ["序号", "发现时间", "店铺", "商品", "SKU", "平台", "基准价", "异动到手价", "价差", "价差率", "状态"]
        writer.writerow(headers)

        for i, alert in enumerate(alerts, 1):
            task = task_map.get(alert.task_id)
            sku = sku_map.get(alert.sku_id)

            row = [
                i,
                alert.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
                task.shop_name if task else "[任务已删除]",
                task.product_title if task else "[商品信息缺失]",
                sku.sku_name if sku else "[SKU信息缺失]",
                task.platform if task else "未知",
                alert.base_price,
                alert.alert_price,
                alert.price_diff,
                alert.price_diff_pct,
                alert.status,
            ]
            writer.writerow(row)

        output.seek(0)

        log.info(
            f"异动记录已导出",
            data={
                "exported_by": user.username,
                "record_count": len(alerts),
                "format": "csv",
            }
        )

        filename = f"alerts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )


@router.get("/price-records/{task_id}")
def export_price_records(
    task_id: int,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    session: Session = Depends(get_db_session),
    user: User = Depends(require_operator)
):
    """导出指定任务的价格记录"""
    task = session.query(MonitorTask).filter(MonitorTask.id == task_id).first()
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")

    query = session.query(PriceRecord).filter(PriceRecord.task_id == task_id)

    if start_date:
        query = query.filter(PriceRecord.collected_at >= start_date)
    if end_date:
        query = query.filter(PriceRecord.collected_at <= end_date)

    records = query.order_by(PriceRecord.collected_at.desc()).all()

    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "价格记录"

        headers = ["采集时间", "SKU", "到手价", "截图"]
        ws.append(headers)

        for record in records:
            sku = session.query(SKU).filter(SKU.id == record.sku_id).first()
            row = [
                record.collected_at.strftime("%Y-%m-%d %H:%M:%S"),
                sku.sku_name if sku else "",
                f"¥{record.price:.2f}",
                record.screenshot_path or "",
            ]
            ws.append(row)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"price_records_{task_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    except ImportError:
        import csv

        output = io.StringIO()
        writer = csv.writer(output)

        headers = ["采集时间", "SKU", "到手价"]
        writer.writerow(headers)

        for record in records:
            sku = session.query(SKU).filter(SKU.id == record.sku_id).first()
            row = [
                record.collected_at.strftime("%Y-%m-%d %H:%M:%S"),
                sku.sku_name if sku else "",
                record.price,
            ]
            writer.writerow(row)

        output.seek(0)

        filename = f"price_records_{task_id}_{datetime.now().strftime('%Y%m%d')}.csv"
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
